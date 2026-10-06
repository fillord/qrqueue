"""Employee attendance and administrator-reviewed face enrollment."""
import uuid
import hashlib
import secrets
from io import BytesIO
from datetime import date, datetime, time, timedelta, timezone
from typing import Annotated
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, Header, Request, Response
from openpyxl import Workbook
from pydantic import BaseModel, Field, model_validator
from redis.asyncio import Redis
from sqlalchemy import delete, func, select
from sqlalchemy.ext.asyncio import AsyncSession
from starlette.concurrency import run_in_threadpool

from app.api.deps import current_admin, current_organization_id
from app.clock import utcnow
from app.db import get_db
from app.models.attendance import AttendanceEvent, AttendanceKiosk, Employee, EmployeeWorkSchedule
from app.models.attendance import AttendanceDepartment as Department
from app.models.enums import AuditActorType
from app.models.organization import Organization
from app.models.user import User
from app.redis import get_redis
from app.services.attendance import code_digest, enrollment_token, new_code, phone_session_token, qr_token, record_attendance, verify_enrollment_token, verify_phone_session, verify_qr
from app.services.audit import log_action
from app.services.employee_import import EmployeeImportError, MAX_IMPORT_BYTES, parse_employee_workbook
from app.services.errors import ServiceError
from app.services.face_attendance import capture_descriptor, decode_review_photo, decode_template, encode_review_photo, encode_template, similarity
from app.services.geo import haversine_m
from app.services.rate_limit import client_ip, enforce_rate_limit

router = APIRouter(prefix="/attendance", tags=["attendance"])


class WorkScheduleDay(BaseModel):
    weekday: int = Field(ge=0, le=6)
    starts_at: time
    ends_at: time

    @model_validator(mode="after")
    def valid_interval(self):
        if self.starts_at >= self.ends_at:
            raise ValueError("Shift end must be after shift start")
        return self


class WorkScheduleReplace(BaseModel):
    schedule: list[WorkScheduleDay] = Field(max_length=7)

    @model_validator(mode="after")
    def unique_weekdays(self):
        weekdays = [item.weekday for item in self.schedule]
        if len(weekdays) != len(set(weekdays)):
            raise ValueError("Each weekday may occur only once")
        return self


class EmployeeCreate(BaseModel):
    full_name: str = Field(min_length=2, max_length=200)
    department_id: uuid.UUID
    position: str | None = Field(default=None, max_length=200)
    user_id: uuid.UUID | None = None
    schedule: list[WorkScheduleDay] = Field(default_factory=list, max_length=7)

    @model_validator(mode="after")
    def unique_schedule_weekdays(self):
        weekdays = [item.weekday for item in self.schedule]
        if len(weekdays) != len(set(weekdays)):
            raise ValueError("Each weekday may occur only once")
        return self


class EmployeeUpdate(BaseModel):
    full_name: str | None = Field(default=None, min_length=2, max_length=200)
    department_id: uuid.UUID | None = None
    position: str | None = Field(default=None, max_length=200)
    is_active: bool | None = None


class FaceCapture(BaseModel):
    images: list[Annotated[str, Field(min_length=1000, max_length=2_000_000)]] = Field(min_length=2, max_length=3)


class EnrollCapture(FaceCapture):
    consent_confirmed: bool


class FaceApproval(BaseModel):
    identity_checked: bool


class PhoneCapture(FaceCapture):
    token: str = Field(max_length=100)
    code: str = Field(pattern=r"^(?:\d{4}|\d{10})$")


class PhoneMarkCapture(PhoneCapture):
    latitude: float | None = Field(default=None, ge=-90, le=90, allow_inf_nan=False)
    longitude: float | None = Field(default=None, ge=-180, le=180, allow_inf_nan=False)
    accuracy_m: float | None = Field(default=None, ge=0, le=100_000, allow_inf_nan=False)


class EnrollmentCapture(PhoneCapture):
    consent_confirmed: bool


class AttendanceSettingsUpdate(BaseModel):
    enrollment_enabled: bool | None = None
    enrollment_on_kiosk: bool | None = None
    geo_enabled: bool | None = None
    geo_latitude: float | None = Field(default=None, ge=-90, le=90, allow_inf_nan=False)
    geo_longitude: float | None = Field(default=None, ge=-180, le=180, allow_inf_nan=False)
    geo_radius_m: int | None = Field(default=None, ge=100, le=5000)


class EventCorrection(BaseModel):
    kind: str = Field(pattern=r"^(in|out)$")
    occurred_at: datetime
    reason: str = Field(min_length=5, max_length=500)


class ManualEvent(EventCorrection):
    employee_id: uuid.UUID


class KioskCreate(BaseModel):
    name: str = Field(min_length=2, max_length=120)


class KioskPair(BaseModel):
    code: str = Field(pattern=r"^\d{6}$")


def kiosk_view(kiosk: AttendanceKiosk) -> dict:
    return {"id": kiosk.id, "name": kiosk.name, "pairing_code": kiosk.pairing_code,
            "paired": kiosk.token_digest is not None, "last_seen_at": kiosk.last_seen_at}


async def _new_kiosk_code(db: AsyncSession) -> str:
    for _ in range(30):
        code = f"{secrets.randbelow(1_000_000):06d}"
        if not await db.scalar(select(AttendanceKiosk.id).where(AttendanceKiosk.pairing_code == code)):
            return code
    raise ServiceError("kiosk_code_unavailable", 503)


async def current_kiosk(db: AsyncSession = Depends(get_db), x_kiosk_token: str | None = Header(default=None, alias="X-Kiosk-Token")) -> AttendanceKiosk:
    if not x_kiosk_token or len(x_kiosk_token) != 64:
        raise ServiceError("kiosk_not_paired", 401)
    digest = hashlib.sha256(x_kiosk_token.encode()).hexdigest()
    kiosk = await db.scalar(select(AttendanceKiosk).where(AttendanceKiosk.token_digest == digest, AttendanceKiosk.deleted_at.is_(None)))
    if kiosk is None:
        raise ServiceError("kiosk_not_paired", 401)
    await _active_org(db, kiosk.organization_id)
    return kiosk


def schedule_view(schedule: list[EmployeeWorkSchedule] | tuple = ()) -> list[dict]:
    return [{"weekday": item.weekday, "starts_at": item.starts_at.strftime("%H:%M"),
             "ends_at": item.ends_at.strftime("%H:%M")} for item in sorted(schedule, key=lambda item: item.weekday)]


def employee_view(employee: Employee, department_name: str | None = None,
                  schedule: list[EmployeeWorkSchedule] | tuple = ()) -> dict:
    return {"id": employee.id, "full_name": employee.full_name, "department": department_name or employee.department,
            "department_id": employee.department_id, "position": employee.position, "code_length": employee.code_length,
            "user_id": employee.user_id, "is_active": employee.is_active,
            "telegram_connected": employee.telegram_chat_id is not None,
            "face_enrolled": employee.face_template is not None, "face_pending": employee.pending_face_template is not None,
            "face_review_photo_available": employee.pending_face_photo is not None,
            "pending_face_submitted_at": employee.pending_face_submitted_at, "deleted_at": employee.deleted_at,
            "schedule": schedule_view(schedule)}


def event_view(event: AttendanceEvent, employee: Employee) -> dict:
    return {"id": event.id, "employee_id": employee.id, "employee_name": employee.full_name,
            "kind": event.kind, "source": event.source, "occurred_at": event.occurred_at,
            "corrected_at": event.corrected_at, "correction_reason": event.correction_reason,
            "needs_review": event.needs_review}


async def _resolve_missing_checkout(db: AsyncSession, employee_id: uuid.UUID, checkout_at: datetime) -> None:
    previous = (await db.scalars(select(AttendanceEvent).where(
        AttendanceEvent.employee_id == employee_id, AttendanceEvent.kind == "in",
        AttendanceEvent.needs_review.is_(True), AttendanceEvent.occurred_at < checkout_at,
    ).order_by(AttendanceEvent.occurred_at.desc()).limit(1))).first()
    if previous is not None:
        previous.needs_review = False


async def _employee(db: AsyncSession, organization_id: uuid.UUID, employee_id: uuid.UUID) -> Employee:
    employee = await db.get(Employee, employee_id)
    if employee is None or employee.organization_id != organization_id or employee.deleted_at is not None:
        raise ServiceError("employee_not_found", 404)
    return employee


async def _active_org(db: AsyncSession, organization_id: uuid.UUID) -> Organization:
    org = await db.get(Organization, organization_id)
    if org is None or org.deleted_at is not None or not org.is_active:
        raise ServiceError("organization_unavailable", 404)
    return org


async def _department(db: AsyncSession, organization_id: uuid.UUID, department_id: uuid.UUID) -> Department:
    department = await db.get(Department, department_id)
    if department is None or department.organization_id != organization_id:
        raise ServiceError("attendance_department_invalid", 404)
    if not department.is_active:
        raise ServiceError("attendance_department_invalid", 422)
    return department


async def _available_code(db: AsyncSession, organization_id: uuid.UUID, used: set[str] | None = None) -> tuple[str, str]:
    if used is None:
        used = set((await db.scalars(select(Employee.code_digest).where(Employee.organization_id == organization_id))).all())
    start = int(new_code())
    for offset in range(10_000):
        code = f"{(start + offset) % 10_000:04d}"
        digest = code_digest(organization_id, code)
        if digest not in used:
            used.add(digest)
            return code, digest
    raise ServiceError("attendance_codes_exhausted", 409)


@router.get("/admin/employees")
async def list_employees(db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> list[dict]:
    rows = (await db.execute(select(Employee, Department.name).outerjoin(Department, Employee.department_id == Department.id).where(
        Employee.organization_id == organization_id, Employee.deleted_at.is_(None)).order_by(Employee.full_name))).all()
    schedule_rows = (await db.scalars(select(EmployeeWorkSchedule).where(
        EmployeeWorkSchedule.organization_id == organization_id))).all()
    schedules: dict[uuid.UUID, list[EmployeeWorkSchedule]] = {}
    for item in schedule_rows:
        schedules.setdefault(item.employee_id, []).append(item)
    return [employee_view(employee, department_name, schedules.get(employee.id, [])) for employee, department_name in rows]


@router.post("/admin/employees", status_code=201)
async def create_employee(payload: EmployeeCreate, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    await _active_org(db, organization_id)
    if len(payload.full_name.strip()) < 2:
        raise ServiceError("employee_name_invalid", 422)
    await db.execute(select(Organization.id).where(Organization.id == organization_id).with_for_update())
    department = await _department(db, organization_id, payload.department_id)
    if payload.user_id:
        linked = await db.get(User, payload.user_id)
        if linked is None or linked.organization_id != organization_id or linked.deleted_at is not None:
            raise ServiceError("employee_user_invalid", 422)
        if await db.scalar(select(Employee.id).where(Employee.user_id == payload.user_id)):
            raise ServiceError("employee_user_already_linked", 409)
    code, digest = await _available_code(db, organization_id)
    employee = Employee(organization_id=organization_id, full_name=payload.full_name.strip(), department=department.name,
                        department_id=department.id, position=payload.position.strip() or None if payload.position else None,
                        user_id=payload.user_id, code_digest=digest, code_length=4)
    db.add(employee)
    await db.flush()
    work_schedule = [EmployeeWorkSchedule(organization_id=organization_id, employee_id=employee.id,
        weekday=item.weekday, starts_at=item.starts_at, ends_at=item.ends_at) for item in payload.schedule]
    db.add_all(work_schedule)
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.employee_created", entity_type="attendance_employee", entity_id=employee.id, organization_id=organization_id)
    await db.commit()
    return {**employee_view(employee, department.name, work_schedule), "code": code}


@router.put("/admin/employees/{employee_id}/schedule")
async def replace_employee_schedule(employee_id: uuid.UUID, payload: WorkScheduleReplace,
                                    db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin),
                                    organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    employee = await _employee(db, organization_id, employee_id)
    await db.execute(delete(EmployeeWorkSchedule).where(
        EmployeeWorkSchedule.organization_id == organization_id,
        EmployeeWorkSchedule.employee_id == employee.id,
    ))
    schedule = [EmployeeWorkSchedule(organization_id=organization_id, employee_id=employee.id,
        weekday=item.weekday, starts_at=item.starts_at, ends_at=item.ends_at) for item in payload.schedule]
    db.add_all(schedule)
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id,
                     action="attendance.schedule_updated", entity_type="attendance_employee",
                     entity_id=employee.id, organization_id=organization_id,
                     payload={"days": len(schedule)})
    await db.commit()
    return employee_view(employee, schedule=schedule)


@router.patch("/admin/employees/{employee_id}")
async def update_employee(employee_id: uuid.UUID, payload: EmployeeUpdate, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    await db.execute(select(Organization.id).where(Organization.id == organization_id).with_for_update())
    employee = await _employee(db, organization_id, employee_id)
    if "department_id" in payload.model_fields_set:
        if payload.department_id is None:
            raise ServiceError("attendance_department_invalid", 422)
        department = await _department(db, organization_id, payload.department_id)
        employee.department_id = department.id
        employee.department = department.name
    for key, value in payload.model_dump(exclude_unset=True).items():
        if key == "department_id":
            continue
        if key == "full_name" and value is not None:
            value = value.strip()
            if len(value) < 2:
                raise ServiceError("employee_name_invalid", 422)
        if key == "position" and value is not None:
            value = value.strip() or None
        setattr(employee, key, value)
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.employee_updated", entity_type="attendance_employee", entity_id=employee.id, organization_id=organization_id)
    await db.commit()
    return employee_view(employee)


@router.delete("/admin/employees/{employee_id}", status_code=204)
async def archive_employee(employee_id: uuid.UUID, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> None:
    employee = await _employee(db, organization_id, employee_id)
    employee.deleted_at = utcnow()
    employee.is_active = False
    employee.face_template = None
    employee.face_consent_at = None
    employee.pending_face_template = None
    employee.pending_face_photo = None
    employee.pending_face_consent_at = None
    employee.pending_face_submitted_at = None
    employee.code_digest = uuid.uuid4().hex + uuid.uuid4().hex
    employee.department_id = None
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.employee_archived", entity_type="attendance_employee", entity_id=employee.id, organization_id=organization_id)
    await db.commit()


@router.post("/admin/employees/{employee_id}/reset-code")
async def reset_code(employee_id: uuid.UUID, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    employee = await _employee(db, organization_id, employee_id)
    await db.execute(select(Organization.id).where(Organization.id == organization_id).with_for_update())
    code, employee.code_digest = await _available_code(db, organization_id)
    employee.code_length = 4
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.code_reset", entity_type="attendance_employee", entity_id=employee.id, organization_id=organization_id)
    await db.commit()
    return {"code": code}


@router.post("/admin/employees/{employee_id}/face")
async def enroll_face(employee_id: uuid.UUID, payload: EnrollCapture, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    employee = await _employee(db, organization_id, employee_id)
    if not payload.consent_confirmed:
        raise ServiceError("face_consent_required", 422)
    vector = await run_in_threadpool(capture_descriptor, payload.images)
    employee.face_template = encode_template(vector)
    employee.face_consent_at = utcnow()
    employee.pending_face_template = None
    employee.pending_face_photo = None
    employee.pending_face_consent_at = None
    employee.pending_face_submitted_at = None
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.face_enrolled", entity_type="attendance_employee", entity_id=employee.id, organization_id=organization_id)
    await db.commit()
    return employee_view(employee)


@router.delete("/admin/employees/{employee_id}/face", status_code=204)
async def remove_face(employee_id: uuid.UUID, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> None:
    employee = await _employee(db, organization_id, employee_id)
    employee.face_template = None
    employee.face_consent_at = None
    employee.pending_face_template = None
    employee.pending_face_photo = None
    employee.pending_face_consent_at = None
    employee.pending_face_submitted_at = None
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.face_removed", entity_type="attendance_employee", entity_id=employee.id, organization_id=organization_id)
    await db.commit()


@router.get("/admin/employees/{employee_id}/pending-face-photo")
async def pending_face_photo(employee_id: uuid.UUID, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> Response:
    employee = await _employee(db, organization_id, employee_id)
    if employee.pending_face_template is None or employee.pending_face_photo is None:
        raise ServiceError("face_review_photo_missing", 404)
    return Response(content=decode_review_photo(employee.pending_face_photo), media_type="image/jpeg",
                    headers={"Cache-Control": "no-store, private", "X-Content-Type-Options": "nosniff"})


@router.post("/admin/employees/{employee_id}/approve-face")
async def approve_face(employee_id: uuid.UUID, payload: FaceApproval, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    employee = await _employee(db, organization_id, employee_id)
    await db.execute(select(Employee.id).where(Employee.id == employee.id).with_for_update())
    if employee.pending_face_template is None or employee.pending_face_consent_at is None:
        raise ServiceError("face_request_missing", 409)
    if employee.pending_face_photo is None:
        raise ServiceError("face_review_photo_missing", 409)
    if not payload.identity_checked:
        raise ServiceError("face_identity_check_required", 422)
    probe = decode_template(employee.pending_face_template)
    others = (await db.scalars(select(Employee).where(Employee.organization_id == organization_id,
        Employee.id != employee.id, Employee.deleted_at.is_(None), Employee.face_template.is_not(None)))).all()
    if any(similarity(probe, decode_template(other.face_template)) >= 0.52 for other in others):
        raise ServiceError("face_already_registered", 409)
    employee.face_template = employee.pending_face_template
    employee.face_consent_at = employee.pending_face_consent_at
    employee.pending_face_template = None
    employee.pending_face_photo = None
    employee.pending_face_consent_at = None
    employee.pending_face_submitted_at = None
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.face_approved", entity_type="attendance_employee", entity_id=employee.id, organization_id=organization_id)
    await db.commit()
    return employee_view(employee)


@router.post("/admin/employees/{employee_id}/reject-face")
async def reject_face(employee_id: uuid.UUID, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    employee = await _employee(db, organization_id, employee_id)
    employee.pending_face_template = None
    employee.pending_face_photo = None
    employee.pending_face_consent_at = None
    employee.pending_face_submitted_at = None
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.face_rejected", entity_type="attendance_employee", entity_id=employee.id, organization_id=organization_id)
    await db.commit()
    return employee_view(employee)


@router.get("/admin/employees/template")
async def employee_template(db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> Response:
    departments = (await db.scalars(select(Department).where(Department.organization_id == organization_id, Department.is_active.is_(True)).order_by(Department.name))).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Сотрудники"
    sheet.append(["ФИО", "Отделение", "Должность"])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 30
    sheet.column_dimensions["C"].width = 30
    from openpyxl.styles import Font, PatternFill
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176E5E")
    choices = workbook.create_sheet("Отделения")
    choices.append(["Допустимые названия отделений"])
    choices.column_dimensions["A"].width = 38
    for department in departments:
        choices.append([department.name])
        # A user-entered department name must remain text, not an Excel formula.
        choices.cell(choices.max_row, 1).data_type = "s"
    output = BytesIO()
    workbook.save(output)
    return Response(content=output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": "attachment; filename=attendance-employees.xlsx"})


@router.post("/admin/employees/import")
async def import_employees(request: Request, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    content_length = request.headers.get("content-length")
    if content_length is not None and (not content_length.isdigit() or int(content_length) > MAX_IMPORT_BYTES):
        raise ServiceError("excel_too_large", 413)
    data = bytearray()
    async for part in request.stream():
        data.extend(part)
        if len(data) > MAX_IMPORT_BYTES:
            raise ServiceError("excel_too_large", 413)
    try:
        rows = await run_in_threadpool(parse_employee_workbook, bytes(data))
    except EmployeeImportError as exc:
        raise ServiceError(exc.code, 413 if exc.code == "excel_too_large" else 422, row=exc.row) from exc
    await db.execute(select(Organization.id).where(Organization.id == organization_id).with_for_update())
    departments = (await db.scalars(select(Department).where(Department.organization_id == organization_id, Department.is_active.is_(True)))).all()
    by_name = {department.name.strip().casefold(): department for department in departments}
    existing = (await db.scalars(select(Employee).where(Employee.organization_id == organization_id, Employee.deleted_at.is_(None)))).all()
    seen = {(employee.full_name.casefold(), employee.department_id) for employee in existing}
    resolved = []
    for row_number, name, department_name, position in rows:
        department = by_name.get(department_name.casefold())
        if department is None:
            raise ServiceError("attendance_department_unknown", 422, row=row_number, department=department_name)
        if (name.casefold(), department.id) in seen:
            raise ServiceError("employee_already_exists", 409, row=row_number)
        seen.add((name.casefold(), department.id))
        resolved.append((name, department, position))
    used = set((await db.scalars(select(Employee.code_digest).where(Employee.organization_id == organization_id))).all())
    created = []
    for name, department, position in resolved:
        code, digest = await _available_code(db, organization_id, used)
        employee = Employee(organization_id=organization_id, full_name=name, department_id=department.id,
                            department=department.name, position=position, code_digest=digest, code_length=4)
        db.add(employee)
        created.append({"full_name": name, "department": department.name, "code": code})
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.employees_imported", entity_type="organization", entity_id=organization_id, organization_id=organization_id, payload={"count": len(created)})
    await db.commit()
    return {"count": len(created), "employees": created}


def _settings_view(org: Organization) -> dict:
    return {"enrollment_enabled": org.attendance_enrollment_enabled,
            "enrollment_on_kiosk": org.attendance_enrollment_on_kiosk,
            "enrollment_token": enrollment_token(org.id, org.attendance_enrollment_version),
            "geo_enabled": org.attendance_geo_enabled,
            "geo_latitude": org.attendance_geo_latitude,
            "geo_longitude": org.attendance_geo_longitude,
            "geo_radius_m": org.attendance_geo_radius_m}


@router.get("/admin/settings")
async def get_attendance_settings(db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    return _settings_view(await _active_org(db, organization_id))


@router.patch("/admin/settings")
async def update_attendance_settings(payload: AttendanceSettingsUpdate, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    org = await _active_org(db, organization_id)
    changes = {key: value for key, value in payload.model_dump(exclude_unset=True).items() if value is not None}
    if changes.get("geo_enabled", org.attendance_geo_enabled) and any(
        changes.get(key, getattr(org, f"attendance_{key}")) is None
        for key in ("geo_latitude", "geo_longitude", "geo_radius_m")
    ):
        raise ServiceError("attendance_geo_config_incomplete", 422)
    for key, value in changes.items():
        setattr(org, f"attendance_{key}", value)
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.settings_updated", entity_type="organization", entity_id=organization_id, organization_id=organization_id,
                     payload={"geo_enabled": org.attendance_geo_enabled} if any(key.startswith("geo_") for key in payload.model_fields_set) else None)
    await db.commit()
    return _settings_view(org)


@router.post("/admin/settings/rotate-enrollment-qr")
async def rotate_enrollment_qr(db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    org = await _active_org(db, organization_id)
    org.attendance_enrollment_version += 1
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.enrollment_qr_rotated", entity_type="organization", entity_id=organization_id, organization_id=organization_id)
    await db.commit()
    return _settings_view(org)


@router.get("/admin/stats")
async def attendance_stats(db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    org = await _active_org(db, organization_id)
    now = utcnow()
    today = now.astimezone(ZoneInfo(org.timezone)).date()
    start = datetime.combine(today - timedelta(days=6), time.min, tzinfo=ZoneInfo(org.timezone))
    employees = (await db.execute(select(Employee.id, Employee.department_id, Department.name, Employee.face_template, Employee.pending_face_template)
        .outerjoin(Department, Employee.department_id == Department.id).where(Employee.organization_id == organization_id,
        Employee.deleted_at.is_(None), Employee.is_active.is_(True)))).all()
    latest = select(AttendanceEvent.employee_id, AttendanceEvent.kind, AttendanceEvent.occurred_at,
        func.row_number().over(partition_by=AttendanceEvent.employee_id, order_by=(AttendanceEvent.occurred_at.desc(), AttendanceEvent.recorded_at.desc())).label("rn"))\
        .where(AttendanceEvent.organization_id == organization_id).subquery()
    last_rows = (await db.execute(select(latest.c.employee_id, latest.c.kind, latest.c.occurred_at).where(latest.c.rn == 1))).all()
    active_ids = {employee_id for employee_id, _, _, _, _ in employees}
    present_ids = {employee_id for employee_id, kind, occurred_at in last_rows if employee_id in active_ids and kind == "in" and now - occurred_at < timedelta(hours=20)}
    recent = (await db.execute(select(AttendanceEvent.kind, AttendanceEvent.occurred_at, AttendanceEvent.employee_id)
        .where(AttendanceEvent.organization_id == organization_id, AttendanceEvent.occurred_at >= start))).all()
    daily = {str(today - timedelta(days=offset)): {"in": 0, "out": 0} for offset in range(6, -1, -1)}
    today_in, today_out = set(), set()
    for kind, occurred_at, employee_id in recent:
        day = str(occurred_at.astimezone(ZoneInfo(org.timezone)).date())
        if day in daily:
            daily[day][kind] += 1
        if day == str(today):
            (today_in if kind == "in" else today_out).add(employee_id)
    by_department = {}
    for employee_id, department_id, department_name, _, _ in employees:
        key = str(department_id) if department_id else "unassigned"
        group = by_department.setdefault(key, {"department_id": key, "name": department_name or "Без отделения", "employees": 0, "present": 0})
        group["employees"] += 1
        group["present"] += int(employee_id in present_ids)
    unresolved = await db.scalar(select(func.count()).select_from(AttendanceEvent).where(AttendanceEvent.organization_id == organization_id, AttendanceEvent.needs_review.is_(True)))
    return {"employees": len(employees), "enrolled": sum(face is not None for _, _, _, face, _ in employees),
            "pending": sum(face is not None for _, _, _, _, face in employees), "present": len(present_ids),
            "arrivals_today": len(today_in), "departures_today": len(today_out), "unresolved": unresolved or 0,
            "daily": [{"date": day, **counts} for day, counts in daily.items()], "departments": list(by_department.values())}


@router.get("/admin/report")
async def attendance_report(date_from: date, date_to: date, department_id: uuid.UUID | None = None,
                            db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin),
                            organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    from app.services.workforce import build_report
    org = await _active_org(db, organization_id)
    return await build_report(db, org, date_from, date_to, department_id)


@router.get("/enroll/context")
async def enrollment_context(token: str, db: AsyncSession = Depends(get_db)) -> dict:
    organization_id, version = verify_enrollment_token(token)
    org = await _active_org(db, organization_id)
    if not org.attendance_enrollment_enabled or org.attendance_enrollment_version != version:
        raise ServiceError("enrollment_link_invalid", 404)
    return {"organization_name": org.name}


@router.post("/enroll/submit")
async def submit_enrollment(payload: EnrollmentCapture, request: Request, db: AsyncSession = Depends(get_db), redis: Redis = Depends(get_redis)) -> dict:
    if not payload.consent_confirmed:
        raise ServiceError("face_consent_required", 422)
    await enforce_rate_limit(redis, scope="attendance-enroll-ip", key=client_ip(request), limit=5)
    organization_id, version = verify_enrollment_token(payload.token)
    org = await _active_org(db, organization_id)
    if not org.attendance_enrollment_enabled or org.attendance_enrollment_version != version:
        raise ServiceError("enrollment_link_invalid", 404)
    digest = code_digest(organization_id, payload.code)
    await enforce_rate_limit(redis, scope="attendance-enroll-code", key=digest, limit=3)
    employee = await db.scalar(select(Employee).where(Employee.organization_id == organization_id, Employee.code_digest == digest,
        Employee.deleted_at.is_(None), Employee.is_active.is_(True)).with_for_update())
    if employee is None or employee.face_template is not None or employee.pending_face_template is not None:
        raise ServiceError("enrollment_not_available", 422)
    vector = await run_in_threadpool(capture_descriptor, payload.images)
    review_photo = await run_in_threadpool(encode_review_photo, payload.images[0])
    employee.pending_face_template = encode_template(vector)
    employee.pending_face_photo = review_photo
    employee.pending_face_consent_at = utcnow()
    employee.pending_face_submitted_at = utcnow()
    await log_action(db, actor_type=AuditActorType.system, actor_id=None, action="attendance.face_submitted", entity_type="attendance_employee", entity_id=employee.id, organization_id=organization_id)
    await db.commit()
    return {"status": "pending", "employee_name": employee.full_name}


@router.get("/admin/events")
async def list_events(db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id), day: date | None = None, limit: int = 500) -> list[dict]:
    limit = max(1, min(limit, 500))
    query = select(AttendanceEvent, Employee).join(Employee, AttendanceEvent.employee_id == Employee.id).where(AttendanceEvent.organization_id == organization_id)
    if day is not None:
        org = await _active_org(db, organization_id)
        start = datetime.combine(day, time.min, tzinfo=ZoneInfo(org.timezone))
        end = datetime.combine(day + timedelta(days=1), time.min, tzinfo=ZoneInfo(org.timezone))
        query = query.where(AttendanceEvent.occurred_at >= start, AttendanceEvent.occurred_at < end)
    rows = (await db.execute(query.order_by(AttendanceEvent.occurred_at.desc()).limit(limit))).all()
    return [event_view(event, employee) for event, employee in rows]


@router.patch("/admin/events/{event_id}")
async def correct_event(event_id: uuid.UUID, payload: EventCorrection, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    event = await db.get(AttendanceEvent, event_id)
    if event is None or event.organization_id != organization_id:
        raise ServiceError("attendance_event_not_found", 404)
    if payload.occurred_at.tzinfo is None or payload.occurred_at > utcnow() + timedelta(minutes=5) or len(payload.reason.strip()) < 5:
        raise ServiceError("attendance_time_invalid", 422)
    employee = await db.get(Employee, event.employee_id)
    old = {"kind": event.kind, "occurred_at": event.occurred_at.isoformat()}
    event.kind = payload.kind
    event.occurred_at = payload.occurred_at
    event.corrected_at = utcnow()
    event.correction_reason = payload.reason.strip()
    event.corrected_by = actor.id
    if event.kind == "out":
        event.needs_review = False
    if event.kind == "out":
        await _resolve_missing_checkout(db, event.employee_id, event.occurred_at)
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.event_corrected", entity_type="attendance_event", entity_id=event.id, organization_id=organization_id, payload={"full_name": employee.full_name, "before": old, "after": {"kind": event.kind, "occurred_at": event.occurred_at.isoformat()}, "reason": event.correction_reason})
    await db.commit()
    return event_view(event, employee)


@router.post("/admin/events", status_code=201)
async def add_manual_event(payload: ManualEvent, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    employee = await _employee(db, organization_id, payload.employee_id)
    if payload.occurred_at.tzinfo is None or payload.occurred_at > utcnow() + timedelta(minutes=5) or len(payload.reason.strip()) < 5:
        raise ServiceError("attendance_time_invalid", 422)
    event = AttendanceEvent(organization_id=organization_id, employee_id=employee.id, kind=payload.kind, source="manual",
                            occurred_at=payload.occurred_at, corrected_at=utcnow(), correction_reason=payload.reason.strip(), corrected_by=actor.id)
    db.add(event)
    await db.flush()
    if event.kind == "out":
        await _resolve_missing_checkout(db, employee.id, event.occurred_at)
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.manual_mark", entity_type="attendance_event", entity_id=event.id, organization_id=organization_id, payload={"full_name": employee.full_name, "kind": event.kind, "reason": event.correction_reason})
    await db.commit()
    return event_view(event, employee)


@router.get("/admin/kiosks")
async def list_kiosks(db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> list[dict]:
    kiosks = (await db.scalars(select(AttendanceKiosk).where(
        AttendanceKiosk.organization_id == organization_id, AttendanceKiosk.deleted_at.is_(None)
    ).order_by(AttendanceKiosk.created_at))).all()
    return [kiosk_view(kiosk) for kiosk in kiosks]


@router.post("/admin/kiosks", status_code=201)
async def create_kiosk(payload: KioskCreate, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    await _active_org(db, organization_id)
    name = payload.name.strip()
    if len(name) < 2:
        raise ServiceError("kiosk_name_invalid", 422)
    kiosk = AttendanceKiosk(organization_id=organization_id, name=name, pairing_code=await _new_kiosk_code(db))
    db.add(kiosk)
    await db.flush()
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.kiosk_created", entity_type="attendance_kiosk", entity_id=kiosk.id, organization_id=organization_id, payload={"name": name})
    await db.commit()
    return kiosk_view(kiosk)


@router.post("/admin/kiosks/{kiosk_id}/unpair")
async def unpair_kiosk(kiosk_id: uuid.UUID, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> dict:
    kiosk = await db.scalar(select(AttendanceKiosk).where(AttendanceKiosk.id == kiosk_id, AttendanceKiosk.organization_id == organization_id, AttendanceKiosk.deleted_at.is_(None)).with_for_update())
    if kiosk is None:
        raise ServiceError("kiosk_not_found", 404)
    kiosk.token_digest = None
    kiosk.last_seen_at = None
    kiosk.pairing_code = await _new_kiosk_code(db)
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.kiosk_unpaired", entity_type="attendance_kiosk", entity_id=kiosk.id, organization_id=organization_id, payload={"name": kiosk.name})
    await db.commit()
    return kiosk_view(kiosk)


@router.delete("/admin/kiosks/{kiosk_id}", status_code=204)
async def archive_kiosk(kiosk_id: uuid.UUID, db: AsyncSession = Depends(get_db), actor: User = Depends(current_admin), organization_id: uuid.UUID = Depends(current_organization_id)) -> None:
    kiosk = await db.scalar(select(AttendanceKiosk).where(AttendanceKiosk.id == kiosk_id, AttendanceKiosk.organization_id == organization_id, AttendanceKiosk.deleted_at.is_(None)).with_for_update())
    if kiosk is None:
        raise ServiceError("kiosk_not_found", 404)
    kiosk.deleted_at = utcnow()
    kiosk.token_digest = None
    kiosk.pairing_code = None
    await log_action(db, actor_type=AuditActorType.user, actor_id=actor.id, action="attendance.kiosk_archived", entity_type="attendance_kiosk", entity_id=kiosk.id, organization_id=organization_id, payload={"name": kiosk.name})
    await db.commit()


@router.post("/kiosk/pair")
async def pair_kiosk(payload: KioskPair, request: Request, db: AsyncSession = Depends(get_db), redis: Redis = Depends(get_redis)) -> dict:
    await enforce_rate_limit(redis, scope="attendance-kiosk-pair", key=client_ip(request), limit=5)
    kiosk = await db.scalar(select(AttendanceKiosk).where(AttendanceKiosk.pairing_code == payload.code, AttendanceKiosk.deleted_at.is_(None)).with_for_update())
    if kiosk is None:
        raise ServiceError("kiosk_pairing_invalid", 404)
    await _active_org(db, kiosk.organization_id)
    token = secrets.token_hex(32)
    kiosk.token_digest = hashlib.sha256(token.encode()).hexdigest()
    kiosk.pairing_code = None
    kiosk.last_seen_at = utcnow()
    await log_action(db, actor_type=AuditActorType.system, actor_id=None, action="attendance.kiosk_paired", entity_type="attendance_kiosk", entity_id=kiosk.id, organization_id=kiosk.organization_id, payload={"name": kiosk.name})
    await db.commit()
    return {"device_token": token, "name": kiosk.name}


@router.get("/kiosk/state")
async def kiosk_state(kiosk: AttendanceKiosk = Depends(current_kiosk), db: AsyncSession = Depends(get_db)) -> dict:
    kiosk.last_seen_at = utcnow()
    await db.commit()
    return {"name": kiosk.name}


@router.get("/kiosk/qr")
async def kiosk_qr(kiosk: AttendanceKiosk = Depends(current_kiosk)) -> dict:
    return {"token": qr_token(kiosk.organization_id), "expires_in": 45}


@router.get("/kiosk/enrollment-qr")
async def kiosk_enrollment_qr(db: AsyncSession = Depends(get_db), kiosk: AttendanceKiosk = Depends(current_kiosk)) -> dict:
    org = await _active_org(db, kiosk.organization_id)
    if not org.attendance_enrollment_enabled or not org.attendance_enrollment_on_kiosk:
        return {"token": None}
    return {"token": enrollment_token(org.id, org.attendance_enrollment_version)}


@router.post("/kiosk/mark")
async def kiosk_mark(payload: FaceCapture, request: Request, db: AsyncSession = Depends(get_db), redis: Redis = Depends(get_redis), kiosk: AttendanceKiosk = Depends(current_kiosk)) -> dict:
    organization_id = kiosk.organization_id
    await enforce_rate_limit(redis, scope="attendance-kiosk", key=f"{kiosk.id}:{client_ip(request)}", limit=12)
    probe = await run_in_threadpool(capture_descriptor, payload.images)
    employees = (await db.scalars(select(Employee).where(Employee.organization_id == organization_id, Employee.deleted_at.is_(None), Employee.is_active.is_(True), Employee.face_template.is_not(None)))).all()
    scored = sorted(((similarity(probe, decode_template(employee.face_template)), employee) for employee in employees), key=lambda item: item[0], reverse=True)
    if not scored or scored[0][0] < 0.60 or (len(scored) > 1 and scored[0][0] - scored[1][0] < 0.08):
        raise ServiceError("face_not_recognized", 422)
    employee = scored[0][1]
    event = await record_attendance(db, employee, "kiosk")
    kiosk.last_seen_at = utcnow()
    await log_action(db, actor_type=AuditActorType.system, actor_id=None, action=f"attendance.marked_{event.kind}", entity_type="attendance_event", entity_id=event.id, organization_id=organization_id, payload={"full_name": employee.full_name, "source": "kiosk", "kiosk": kiosk.name})
    await db.commit()
    return event_view(event, employee)


@router.get("/phone/context")
async def phone_context(token: str, db: AsyncSession = Depends(get_db)) -> dict:
    organization_id = verify_qr(token)
    org = await _active_org(db, organization_id)
    return {"organization_name": org.name, "session_token": phone_session_token(organization_id),
            "geo_required": org.attendance_geo_enabled}


@router.post("/phone/mark")
async def phone_mark(payload: PhoneMarkCapture, request: Request, db: AsyncSession = Depends(get_db), redis: Redis = Depends(get_redis)) -> dict:
    await enforce_rate_limit(redis, scope="attendance-phone-ip", key=client_ip(request), limit=8)
    organization_id = verify_phone_session(payload.token)
    org = await _active_org(db, organization_id)
    if org.attendance_geo_enabled:
        if payload.latitude is None or payload.longitude is None or payload.accuracy_m is None:
            raise ServiceError("attendance_geo_required", 422)
        if org.attendance_geo_latitude is None or org.attendance_geo_longitude is None or org.attendance_geo_radius_m is None:
            raise ServiceError("attendance_geo_config_incomplete", 503)
        if payload.accuracy_m > org.attendance_geo_radius_m:
            raise ServiceError("attendance_geo_inaccurate", 422)
        if haversine_m(org.attendance_geo_latitude, org.attendance_geo_longitude, payload.latitude, payload.longitude) > org.attendance_geo_radius_m:
            raise ServiceError("attendance_geo_out_of_range", 403)
    digest = code_digest(organization_id, payload.code)
    await enforce_rate_limit(redis, scope="attendance-phone-code", key=digest, limit=5)
    employee = await db.scalar(select(Employee).where(Employee.organization_id == organization_id, Employee.code_digest == digest, Employee.deleted_at.is_(None), Employee.is_active.is_(True)))
    if employee is None or employee.face_template is None or employee.face_consent_at is None:
        raise ServiceError("face_not_recognized", 422)
    probe = await run_in_threadpool(capture_descriptor, payload.images)
    if similarity(probe, decode_template(employee.face_template)) < 0.52:
        raise ServiceError("face_not_recognized", 422)
    event = await record_attendance(db, employee, "phone")
    await log_action(db, actor_type=AuditActorType.system, actor_id=None, action=f"attendance.marked_{event.kind}", entity_type="attendance_event", entity_id=event.id, organization_id=organization_id, payload={"full_name": employee.full_name, "source": "phone"})
    await db.commit()
    return event_view(event, employee)
