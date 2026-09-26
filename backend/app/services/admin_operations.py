import uuid
from datetime import date, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.clock import utcnow
from app.models.cabinet import Cabinet
from app.models.enums import CabinetStatus, QueueStatus, TicketStatus
from app.models.organization import Organization
from app.models.queue import Queue
from app.models.ticket import Ticket
from app.models.tv_screen import TVScreen
from app.services.analytics import _range_to_utc, get_analytics
from app.services.queues import list_queues_with_waiting_counts

LONG_WAIT_MINUTES = 20
SCREEN_OFFLINE_SECONDS = 90


async def get_admin_problems(db: AsyncSession, organization_id: uuid.UUID) -> dict:
    now = utcnow()
    pairs = await list_queues_with_waiting_counts(db, organization_id)
    queues = [(queue, count) for queue, count in pairs if queue.is_active and count > 0]
    queue_ids = [queue.id for queue, _ in queues]
    oldest_by_queue = {}
    staffed_queue_ids = set()
    if queue_ids:
        oldest_by_queue = dict((await db.execute(
            select(Ticket.queue_id, func.min(Ticket.created_at))
            .where(Ticket.organization_id == organization_id, Ticket.queue_id.in_(queue_ids),
                   Ticket.status == TicketStatus.waiting)
            .group_by(Ticket.queue_id)
        )).all())
        staffed_queue_ids = set((await db.scalars(
            select(Cabinet.queue_id).where(
                Cabinet.organization_id == organization_id, Cabinet.queue_id.in_(queue_ids),
                Cabinet.deleted_at.is_(None), Cabinet.is_active.is_(True),
                Cabinet.status.in_((CabinetStatus.free, CabinetStatus.busy)),
            ).distinct()
        )).all())

    items = []
    for queue, count in queues:
        oldest = oldest_by_queue.get(queue.id)
        wait_minutes = max(0, int((now - oldest).total_seconds() // 60)) if oldest else None
        code = None
        if queue.status != QueueStatus.open:
            code = "queue_not_open"
        elif queue.id not in staffed_queue_ids:
            code = "no_cabinet"
        elif wait_minutes is not None and wait_minutes >= LONG_WAIT_MINUTES:
            code = "long_wait"
        if code:
            items.append({
                "code": code, "severity": "warning" if code == "long_wait" else "critical",
                "name": queue.name, "queue_id": queue.id, "waiting_count": count,
                "wait_minutes": wait_minutes,
            })

    screens = (await db.scalars(select(TVScreen).where(
        TVScreen.organization_id == organization_id, TVScreen.device_token.is_not(None),
    ))).all()
    cutoff = now - timedelta(seconds=SCREEN_OFFLINE_SECONDS)
    for screen in screens:
        if screen.last_seen_at is None or screen.last_seen_at < cutoff:
            items.append({
                "code": "screen_offline", "severity": "warning",
                "name": screen.name, "screen_id": screen.id,
            })
    items.sort(key=lambda item: (item["severity"] != "critical", item["name"].casefold()))
    return {"generated_at": now, "items": items}


async def get_daily_report(db: AsyncSession, organization: Organization, day: date | None) -> dict:
    day = day or utcnow().astimezone(ZoneInfo(organization.timezone)).date()
    start, end = _range_to_utc(organization.timezone, day, day)
    wait_seconds = func.extract("epoch", Ticket.called_at - Ticket.created_at)
    rows = (await db.execute(
        select(
            Queue.id, Queue.name, func.count(),
            func.count().filter(Ticket.status == TicketStatus.served),
            func.count().filter(Ticket.status == TicketStatus.no_show),
            func.count().filter(Ticket.status == TicketStatus.left),
            func.count().filter(Ticket.status.in_((TicketStatus.waiting, TicketStatus.called,
                                                 TicketStatus.confirmed, TicketStatus.serving))),
            func.avg(wait_seconds).filter(Ticket.status == TicketStatus.served,
                                          Ticket.called_at.is_not(None)),
        )
        .select_from(Ticket).join(Queue, Queue.id == Ticket.queue_id)
        .where(Ticket.organization_id == organization.id, Ticket.created_at >= start,
               Ticket.created_at < end)
        .group_by(Queue.id, Queue.name).order_by(Queue.name)
    )).all()
    by_queue = [{
        "queue_id": queue_id, "name": name, "issued_count": issued,
        "served_count": served, "no_show_count": no_show, "left_count": left,
        "active_count": waiting, "avg_wait_seconds": round(avg_wait) if avg_wait is not None else None,
    } for queue_id, name, issued, served, no_show, left, waiting, avg_wait in rows]
    analytics = await get_analytics(db, organization_id=organization.id,
                                   timezone_name=organization.timezone, date_from=day, date_to=day)
    return {
        "day": day, "organization_name": organization.name, "timezone": organization.timezone,
        "issued_count": sum(row["issued_count"] for row in by_queue),
        "served_count": analytics["served_count"], "no_show_count": analytics["no_show_count"],
        "left_count": analytics["left_count"],
        "active_count": sum(row["active_count"] for row in by_queue),
        "avg_wait_seconds": analytics["avg_wait_seconds"],
        "avg_serving_seconds": analytics["avg_serving_seconds"],
        "by_queue": by_queue,
        "by_operator": [{"operator_id": row["operator_id"], "full_name": row["full_name"],
                         "served_count": row["served_count"]} for row in analytics["by_operator"]],
    }


_REPORT_LABELS = {
    "ru": ("Отчёт", "Дневной отчёт", "Дата", "Часовой пояс", "Талонов создано", "Обслужено",
           "Неявки", "Ушли", "Ещё активны", "Среднее ожидание, мин", "Средний приём, мин",
           "Очередь", "Создано", "Активны", "Сотрудник"),
    "en": ("Report", "Daily report", "Date", "Time zone", "Tickets created", "Served",
           "No-shows", "Left", "Still active", "Average wait, min", "Average service, min",
           "Queue", "Created", "Active", "Employee"),
    "kk": ("Есеп", "Күндік есеп", "Күні", "Уақыт белдеуі", "Құрылған талондар", "Қызмет көрсетілді",
           "Келмегендер", "Шыққандар", "Әлі белсенді", "Орташа күту, мин", "Орташа қабылдау, мин",
           "Кезек", "Құрылған", "Белсенді", "Қызметкер"),
}


def daily_report_xlsx(report: dict, language: str = "ru") -> bytes:
    (sheet_name, title, date_label, timezone_label, issued_label, served_label,
     no_show_label, left_label, active_label, wait_label, service_label,
     queue_label, created_label, active_short_label, operator_label) = _REPORT_LABELS[language]
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name
    sheet.append([title, report["organization_name"]])
    sheet.append([date_label, report["day"].isoformat()])
    sheet.append([timezone_label, report["timezone"]])
    sheet.append([])
    for label, key in ((issued_label, "issued_count"), (served_label, "served_count"),
                       (no_show_label, "no_show_count"), (left_label, "left_count"),
                       (active_label, "active_count")):
        sheet.append([label, report[key]])
    sheet.append([wait_label, round(report["avg_wait_seconds"] / 60, 1) if report["avg_wait_seconds"] is not None else None])
    sheet.append([service_label, round(report["avg_serving_seconds"] / 60, 1) if report["avg_serving_seconds"] is not None else None])
    sheet.append([])
    sheet.append([queue_label, created_label, served_label, no_show_label, left_label,
                  active_short_label, wait_label])
    header_row = sheet.max_row
    for row in report["by_queue"]:
        sheet.append([row["name"], row["issued_count"], row["served_count"], row["no_show_count"],
                      row["left_count"], row["active_count"],
                      round(row["avg_wait_seconds"] / 60, 1) if row["avg_wait_seconds"] is not None else None])
        sheet.cell(sheet.max_row, 1).data_type = "s"
    sheet.append([])
    sheet.append([operator_label, served_label])
    operator_header_row = sheet.max_row
    for row in report["by_operator"]:
        sheet.append([row["full_name"], row["served_count"]])
        sheet.cell(sheet.max_row, 1).data_type = "s"
    sheet["B1"].data_type = "s"
    for row_number in (header_row, operator_header_row):
        for cell in sheet[row_number]:
            cell.fill = PatternFill("solid", fgColor="EAF1FF")
            cell.font = Font(bold=True, color="17233B")
            cell.alignment = Alignment(wrap_text=True)
    for width, column in ((34, "A"), (17, "B"), (17, "C"), (15, "D"), (15, "E"), (15, "F"), (25, "G")):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = f"B{header_row + 1}"
    stream = BytesIO()
    book.save(stream)
    return stream.getvalue()
