"""Strict, all-or-nothing XLSX parsing for the staff directory."""
from io import BytesIO
from zipfile import ZipFile

from openpyxl import load_workbook

MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 10 * 1024 * 1024
MAX_ROWS = 1000
HEADERS = ("фио", "отделение", "должность")


class EmployeeImportError(Exception):
    def __init__(self, code: str, row: int | None = None, department: str | None = None):
        self.code, self.row, self.department = code, row, department


def parse_employee_workbook(data: bytes) -> list[tuple[int, str, str, str | None]]:
    if not data or len(data) > MAX_IMPORT_BYTES:
        raise EmployeeImportError("excel_too_large")
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 100 or sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES:
                raise EmployeeImportError("excel_too_large")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except EmployeeImportError:
        raise
    except Exception as exc:
        raise EmployeeImportError("invalid_excel_format") from exc
    try:
        if "Сотрудники" not in workbook.sheetnames:
            raise EmployeeImportError("invalid_excel_headers")
        sheet = workbook["Сотрудники"]
        header = next(sheet.iter_rows(min_row=1, max_row=1, max_col=3, values_only=True), ())
        if tuple(str(value).strip().casefold() for value in header) != HEADERS:
            raise EmployeeImportError("invalid_excel_headers")
        if sheet.max_row is not None and sheet.max_row > MAX_ROWS + 1:
            raise EmployeeImportError("excel_too_many_rows")
        rows = []
        seen = set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=3), start=2):
            if all(cell.value in (None, "") for cell in cells):
                continue
            if len(rows) >= MAX_ROWS or any(cell.data_type == "f" for cell in cells):
                raise EmployeeImportError("invalid_employee_row", row_number)
            if any(cell.value is not None and not isinstance(cell.value, str) for cell in cells):
                raise EmployeeImportError("invalid_employee_row", row_number)
            name = str(cells[0].value or "").strip()
            department = str(cells[1].value or "").strip()
            position = str(cells[2].value or "").strip() or None
            if not (2 <= len(name) <= 200 and 1 <= len(department) <= 200 and (position is None or len(position) <= 200)):
                raise EmployeeImportError("invalid_employee_row", row_number)
            key = (name.casefold(), department.casefold())
            if key in seen:
                raise EmployeeImportError("duplicate_employee_row", row_number)
            seen.add(key)
            rows.append((row_number, name, department, position))
        if not rows:
            raise EmployeeImportError("empty_employee_file")
        return rows
    finally:
        workbook.close()
