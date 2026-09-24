"""Validate the published weekly schedule XLSX template before changing data."""

from collections import defaultdict
from datetime import time
from io import BytesIO
import re
from zipfile import ZipFile

from openpyxl import load_workbook
from pydantic import TypeAdapter, ValidationError

from app.schemas.tv_signage import ScheduleItemCreate
from app.schemas.validation import Name

HEADERS = ("отделение", "врач", "специализация", "кабинет", "пн", "вт", "ср", "чт", "пт", "сб", "вс")
MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 10 * 1024 * 1024
MAX_SCHEDULE_ROWS = 1000
_NAME = TypeAdapter(Name)
_INTERVAL = re.compile(r"^((?:[01]\d|2[0-3]):[0-5]\d)\s*[-–—]\s*((?:[01]\d|2[0-3]):[0-5]\d)$")


class ScheduleImportError(Exception):
    def __init__(self, code: str, row: int | None = None):
        self.code = code
        self.row = row


def _parse_time(value: str) -> time:
    hour, minute = (int(part) for part in value.split(":"))
    return time(hour, minute)


def _period(value: object, row: int) -> tuple[time, time] | None:
    if value is None or (isinstance(value, str) and value.strip().casefold() in ("", "-", "—", "выходной")):
        return None
    if not isinstance(value, str):
        raise ScheduleImportError("invalid_schedule_row", row)
    match = _INTERVAL.fullmatch(value.strip())
    if not match:
        raise ScheduleImportError("invalid_schedule_row", row)
    start, end = _parse_time(match.group(1)), _parse_time(match.group(2))
    if start >= end:
        raise ScheduleImportError("invalid_schedule_row", row)
    return start, end


def parse_schedule_workbook(data: bytes) -> dict[str, list[ScheduleItemCreate]]:
    if not data or len(data) > MAX_IMPORT_BYTES:
        raise ScheduleImportError("excel_too_large")
    try:
        with ZipFile(BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 100 or sum(member.file_size for member in members) > MAX_UNCOMPRESSED_BYTES:
                raise ScheduleImportError("excel_too_large")
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except ScheduleImportError:
        raise
    except Exception as exc:
        raise ScheduleImportError("invalid_excel_format") from exc

    try:
        if "Расписание" not in workbook.sheetnames:
            raise ScheduleImportError("invalid_excel_headers")
        sheet = workbook["Расписание"]
        header = next(sheet.iter_rows(min_row=1, max_row=1, max_col=11, values_only=True), ())
        normalized = tuple(value.strip().casefold() if isinstance(value, str) else value for value in header)
        if normalized != HEADERS:
            raise ScheduleImportError("invalid_excel_headers")
        if sheet.max_row is not None and sheet.max_row > MAX_SCHEDULE_ROWS + 1:
            raise ScheduleImportError("excel_too_many_rows")

        grouped: dict[str, list[ScheduleItemCreate]] = defaultdict(list)
        row_count = 0
        seen_doctors = set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=11), start=2):
            values = [cell.value for cell in cells]
            if all(value is None or value == "" for value in values):
                continue
            row_count += 1
            if row_count > MAX_SCHEDULE_ROWS or any(cell.data_type == "f" for cell in cells):
                raise ScheduleImportError("invalid_schedule_row", row_number)
            try:
                department = _NAME.validate_python(values[0])
                doctor = _NAME.validate_python(values[1])
                if (department, doctor) in seen_doctors:
                    raise ScheduleImportError("duplicate_schedule_row", row_number)
                service = None if values[2] in (None, "") else str(values[2]).strip() or None
                room = None if values[3] in (None, "") else str(values[3]).strip() or None
                day_periods = [_period(value, row_number) for value in values[4:]]
                if not any(day_periods):
                    raise ScheduleImportError("invalid_schedule_row", row_number)
                for weekday, period in enumerate(day_periods):
                    if period is None:
                        continue
                    start, end = period
                    item = ScheduleItemCreate.model_validate({
                        "doctor_name": doctor, "service_name": service, "room": room,
                        "weekday": weekday, "starts_at": start, "ends_at": end,
                        "sort_order": row_number - 2,
                    })
                    grouped[department].append(item)
                seen_doctors.add((department, doctor))
            except (ValidationError, ValueError, TypeError) as exc:
                raise ScheduleImportError("invalid_schedule_row", row_number) from exc
        if not grouped:
            raise ScheduleImportError("empty_schedule_file")
        return dict(grouped)
    finally:
        workbook.close()
