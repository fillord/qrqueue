"""One calendar/report calculation shared by the UI, payroll export and reminders."""
import calendar
import re
import uuid
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.clock import utcnow
from app.models.attendance import AttendanceEvent, Employee, EmployeeCalendarDay, EmployeeWorkSchedule
from app.models.department import Department
from app.models.organization import Organization
from app.services.errors import ServiceError


def month_range(month: str) -> tuple[date, date]:
    if not re.fullmatch(r"\d{4}-\d{2}", month):
        raise ServiceError("attendance_month_invalid", 422)
    try:
        year, number = map(int, month.split("-"))
        start = date(year, number, 1)
        end = date(year, number, calendar.monthrange(year, number)[1])
        if year < 2000 or year > 2100:
            raise ValueError()
        return start, end
    except ValueError as exc:
        raise ServiceError("attendance_month_invalid", 422) from exc


def worked_intervals(events: list[AttendanceEvent]) -> tuple[int, bool]:
    """Count only closed in/out intervals; never count breaks or invent a checkout."""
    opened = None
    minutes = 0
    invalid = False
    for event in events:
        if event.kind == "in":
            if opened is not None:
                invalid = True
            else:
                opened = event.occurred_at
        elif opened is None:
            invalid = True
        else:
            minutes += max(0, int((event.occurred_at - opened).total_seconds() // 60))
            opened = None
    return minutes, invalid or opened is not None


async def build_report(db: AsyncSession, org: Organization, date_from: date, date_to: date,
                       department_id: uuid.UUID | None = None, *, all_days: bool = False) -> dict:
    if date_to < date_from or (date_to - date_from).days > 92:
        raise ServiceError("attendance_report_range_invalid", 422)
    tz = ZoneInfo(org.timezone)
    start = datetime.combine(date_from, time.min, tzinfo=tz).astimezone(timezone.utc)
    query = select(Employee, Department.name).outerjoin(Department, Employee.department_id == Department.id).where(
        Employee.organization_id == org.id, or_(Employee.deleted_at.is_(None), Employee.deleted_at >= start))
    # Include inactive staff with historical marks: payroll must not lose their work.
    if department_id is not None:
        department = await db.get(Department, department_id)
        if department is None or department.organization_id != org.id:
            raise ServiceError("attendance_department_invalid", 404)
        query = query.where(Employee.department_id == department_id)
    people = (await db.execute(query.order_by(Employee.full_name))).all()
    ids = [person.id for person, _ in people]
    schedules = (await db.scalars(select(EmployeeWorkSchedule).where(EmployeeWorkSchedule.employee_id.in_(ids),
        EmployeeWorkSchedule.organization_id == org.id))).all()
    weekly = {(item.employee_id, item.weekday): item for item in schedules}
    overrides = (await db.scalars(select(EmployeeCalendarDay).where(EmployeeCalendarDay.employee_id.in_(ids),
        EmployeeCalendarDay.organization_id == org.id, EmployeeCalendarDay.day.between(date_from, date_to)))).all()
    dated = {(item.employee_id, item.day): item for item in overrides}
    end = datetime.combine(date_to + timedelta(days=1), time.min, tzinfo=tz).astimezone(timezone.utc)
    events = (await db.scalars(select(AttendanceEvent).where(AttendanceEvent.organization_id == org.id,
        AttendanceEvent.employee_id.in_(ids), AttendanceEvent.occurred_at >= start,
        AttendanceEvent.occurred_at < end).order_by(AttendanceEvent.occurred_at, AttendanceEvent.recorded_at))).all()
    by_day = {}
    for event in events:
        by_day.setdefault((event.employee_id, event.occurred_at.astimezone(tz).date()), []).append(event)
    now = utcnow()
    rows = []
    totals = dict(scheduled=0, completed=0, absent=0, late=0, early_leave=0, incomplete=0, overtime_minutes=0)
    day = date_from
    while day <= date_to:
        for employee, department in people:
            override = dated.get((employee.id, day))
            shift = override if override is not None else weekly.get((employee.id, day.weekday()))
            kind = override.kind if override is not None else "shift" if shift else "off"
            marks = by_day.get((employee.id, day), [])
            if not employee.is_active and not marks and override is None:
                kind, shift = "off", None
            if not all_days and shift is None and not marks:
                continue
            planned_start = datetime.combine(day, shift.starts_at, tzinfo=tz) if kind == "shift" else None
            planned_end = datetime.combine(day, shift.ends_at, tzinfo=tz) if kind == "shift" else None
            arrivals = [event.occurred_at for event in marks if event.kind == "in"]
            departures = [event.occurred_at for event in marks if event.kind == "out"]
            first_in = min(arrivals) if arrivals else None
            last_out = max(departures) if departures else None
            minutes, invalid = worked_intervals(marks)
            invalid = invalid or any(event.needs_review for event in marks)
            late = max(0, int((first_in - planned_start).total_seconds() // 60)) if first_in and planned_start else 0
            early = max(0, int((planned_end - last_out).total_seconds() // 60)) if last_out and planned_end else 0
            planned_minutes = int((planned_end - planned_start).total_seconds() // 60) if planned_start else 0
            overtime = max(0, minutes - planned_minutes)
            if kind != "shift":
                status = kind
            elif first_in and last_out and not invalid:
                status = "late_early" if late and early else "late" if late else "early_leave" if early else "completed"
                totals["completed"] += 1
            elif marks:
                status = "in_progress" if first_in and not last_out and now < planned_end else "incomplete"
                totals["incomplete"] += int(status == "incomplete")
            elif now < planned_end:
                status = "planned"
            else:
                status = "absent"
                totals["absent"] += 1
            totals["scheduled"] += int(kind == "shift")
            totals["late"] += int(late > 0)
            totals["early_leave"] += int(early > 0)
            totals["overtime_minutes"] += overtime
            rows.append(dict(date=day, employee_id=employee.id, employee_name=employee.full_name,
                department=department or employee.department, position=employee.position,
                planned_start=planned_start, planned_end=planned_end, planned_minutes=planned_minutes,
                first_in=first_in, last_out=last_out, worked_minutes=minutes if marks else None,
                late_minutes=late, early_leave_minutes=early, overtime_minutes=overtime, status=status,
                calendar_kind=kind, calendar_id=override.id if override else None,
                reason=override.reason if override else None, needs_review=invalid,
                telegram_connected=employee.telegram_chat_id is not None))
        day += timedelta(days=1)
    rows.sort(key=lambda item: (-item["date"].toordinal(), item["employee_name"].casefold()))
    scheduled_ids = {item.employee_id for item in schedules} | {item.employee_id for item in overrides if item.kind == "shift"}
    return dict(timezone=org.timezone, date_from=date_from, date_to=date_to,
        missing_schedule=sum(person.id not in scheduled_ids for person, _ in people if person.is_active),
        totals=totals, rows=rows)


_EXPORT = {
    "ru": ("Табель", "Детализация", "ФИО", "Отделение", "Должность", "План, ч", "Факт, ч", "Опоздание, мин", "Переработка, ч", "Дата", "Приход", "Уход", "Статус", "Примечание"),
    "kk": ("Табель", "Мәліметтер", "Аты-жөні", "Бөлім", "Лауазым", "Жоспар, сағ", "Нақты, сағ", "Кешігу, мин", "Артық жұмыс, сағ", "Күн", "Келу", "Кету", "Күй", "Ескерту"),
    "en": ("Timesheet", "Details", "Employee", "Department", "Position", "Planned hours", "Worked hours", "Late minutes", "Overtime hours", "Date", "Arrival", "Departure", "Status", "Note"),
}
_CODES = {"off": "OFF", "vacation": "VAC", "sick": "SICK", "absence": "EXC", "absent": "ABS",
          "incomplete": "?", "in_progress": "…", "planned": "PLAN"}


def safe_text(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def timesheet_workbook(report: dict, month: str, language: str = "ru") -> bytes:
    labels = _EXPORT.get(language, _EXPORT["ru"])
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = labels[0]
    start, end = month_range(month)
    days = list(range(1, end.day + 1))
    sheet.append([f"{labels[0]} · {month} · {report['timezone']}"])
    legends = {
        "ru": "OFF — выходной; VAC — отпуск; SICK — больничный; EXC — согласованное отсутствие; ABS — неявка; PLAN — план; … — на работе; ? — проверить отметки. Факт: только закрытые интервалы, без перерывов.",
        "kk": "OFF — демалыс күні; VAC — демалыс; SICK — ауру демалысы; EXC — келісілген келмеу; ABS — келмеді; PLAN — жоспар; … — жұмыста; ? — белгілерді тексеріңіз. Нақты: тек жабық аралықтар, үзіліссіз.",
        "en": "OFF — day off; VAC — vacation; SICK — sick leave; EXC — approved absence; ABS — no-show; PLAN — planned; … — at work; ? — review marks. Worked hours count closed intervals, excluding breaks.",
    }
    sheet.append([legends.get(language, legends["ru"])])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(days) + 7)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(days) + 7)
    sheet["A2"].alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[2].height = 36
    sheet.append([labels[2], labels[3], labels[4], *days, *labels[5:9]])
    by_person = {}
    for row in report["rows"]:
        by_person.setdefault(row["employee_id"], []).append(row)
    for rows in sorted(by_person.values(), key=lambda rows: rows[0]["employee_name"].casefold()):
        person = rows[0]
        by_date = {item["date"].day: item for item in rows}
        values = []
        for day in days:
            item = by_date.get(day)
            if item is None:
                values.append("")
            elif item["needs_review"]:
                values.append("?")
            elif item["worked_minutes"]:
                values.append(round(item["worked_minutes"] / 60, 2))
            else:
                values.append(_CODES.get(item["status"], 0))
        sheet.append([safe_text(person["employee_name"]), safe_text(person["department"]), safe_text(person["position"]),
            *values, round(sum(item["planned_minutes"] for item in rows) / 60, 2),
            round(sum(item["worked_minutes"] or 0 for item in rows) / 60, 2),
            sum(item["late_minutes"] for item in rows), round(sum(item["overtime_minutes"] for item in rows) / 60, 2)])
    sheet.freeze_panes = "D4"
    sheet.auto_filter.ref = f"A3:{get_column_letter(sheet.max_column)}{sheet.max_row}"
    sheet.column_dimensions["A"].width = 30
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 22
    for column in range(4, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 9 if column <= len(days) + 3 else 18
    details = workbook.create_sheet(labels[1])
    details.append([labels[9], labels[2], labels[3], *labels[5:9], labels[10], labels[11], labels[12], labels[13]])
    tz = ZoneInfo(report["timezone"])
    for item in sorted(report["rows"], key=lambda item: (item["date"], item["employee_name"])):
        details.append([item["date"].isoformat(), safe_text(item["employee_name"]), safe_text(item["department"]),
            round(item["planned_minutes"] / 60, 2), round((item["worked_minutes"] or 0) / 60, 2),
            item["late_minutes"], round(item["overtime_minutes"] / 60, 2),
            item["first_in"].astimezone(tz).strftime("%H:%M") if item["first_in"] else "",
            item["last_out"].astimezone(tz).strftime("%H:%M") if item["last_out"] else "",
            item["status"], safe_text(item["reason"])])
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    for column in range(1, 12):
        details.column_dimensions[get_column_letter(column)].width = 24
    for worksheet, header in ((sheet, 3), (details, 1)):
        for cell in worksheet[header]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="174D43")
            cell.alignment = Alignment(wrap_text=True)
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A3
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.print_title_rows = f"1:{header}"
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
