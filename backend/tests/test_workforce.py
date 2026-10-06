import uuid
from datetime import date, datetime, time, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.clock import utcnow
from app.models.attendance import AttendanceEvent, Employee, EmployeeCalendarDay, EmployeeWorkSchedule
from app.models.department import Department
from app.models.enums import UserRole
from app.redis import redis_client
from app.services import telegram
from app.services.workforce import build_report, month_range, worked_intervals
from app.workers.telegram import run_reminders_once
from tests.utils import login


async def setup_employee(client, db_session, make_user, make_organization):
    org = await make_organization(name='Workforce', timezone='UTC')
    admin, password = await make_user(email='workforce@example.com', role=UserRole.org_admin, organization_id=org.id)
    department = Department(organization_id=org.id, name='Clinic')
    db_session.add(department)
    await db_session.commit()
    await login(client, admin.email, password)
    response = await client.post('/api/attendance/admin/employees', json={
        'full_name': '=Dangerous name', 'department_id': str(department.id),
        'schedule': [{'weekday': weekday, 'starts_at': '09:00', 'ends_at': '18:00'} for weekday in range(7)]})
    assert response.status_code == 201, response.text
    employee = await db_session.get(Employee, uuid.UUID(response.json()['id']))
    return org, employee, admin


def test_closed_intervals_exclude_breaks_and_flag_missing_checkout():
    start = datetime(2026, 10, 1, 9, tzinfo=timezone.utc)
    events = [SimpleNamespace(kind=kind, occurred_at=start + timedelta(minutes=minutes))
              for kind, minutes in [('in', 0), ('out', 180), ('in', 240), ('out', 540)]]
    assert worked_intervals(events) == (480, False)
    assert worked_intervals(events[:-1]) == (180, True)
    assert worked_intervals(events[1:]) == (300, True)


@pytest.mark.parametrize('month', ['2026-13', '2026-00', '26-10', '../2026-10', '1900-01', '2101-01'])
def test_month_validation(month):
    from app.services.errors import ServiceError
    with pytest.raises(ServiceError):
        month_range(month)


async def test_calendar_overrides_absence_reset_and_excel(client, db_session, make_user, make_organization):
    org, employee, admin = await setup_employee(client, db_session, make_user, make_organization)
    day = (utcnow() - timedelta(days=2)).date()
    payload = dict(employee_id=str(employee.id), date_from=day.isoformat(), date_to=day.isoformat(), kind='vacation', reason='Approved holiday')
    response = await client.post('/api/attendance/admin/calendar', json=payload)
    assert response.status_code == 200, response.text
    assert response.json()['days'] == 1
    for kind in ('vacation', 'sick', 'absence', 'off'):
        payload['kind'] = kind
        assert (await client.post('/api/attendance/admin/calendar', json=payload)).status_code == 200
        report = await build_report(db_session, org, day, day)
        assert report['totals']['scheduled'] == 0 and report['totals']['absent'] == 0
        assert report['rows'][0]['status'] == kind
    payload.update(kind='shift', starts_at='10:00', ends_at='17:00')
    assert (await client.post('/api/attendance/admin/calendar', json=payload)).status_code == 200
    report = await build_report(db_session, org, day, day)
    assert report['rows'][0]['planned_minutes'] == 420
    row_id = report['rows'][0]['calendar_id']
    month = day.strftime('%Y-%m')
    response = await client.get('/api/attendance/admin/timesheet.xlsx', params={'month': month, 'language': 'ru'})
    assert response.status_code == 200 and response.headers['cache-control'] == 'no-store'
    book = load_workbook(BytesIO(response.content))
    assert book.sheetnames == ['Табель', 'Детализация']
    assert book['Табель']['A4'].value == "'=Dangerous name"
    assert book['Табель']['A4'].data_type != 'f'
    assert (await client.post(f'/api/attendance/admin/calendar/{row_id}/reset', json={'reason': 'Restore default'})).status_code == 200
    report = await build_report(db_session, org, day, day)
    assert report['rows'][0]['planned_minutes'] == 540 and report['rows'][0]['calendar_id'] is None
    assert not (await db_session.scalars(select(EmployeeCalendarDay).where(EmployeeCalendarDay.employee_id == employee.id))).all()
    invalid = dict(payload, ends_at='08:00')
    assert (await client.post('/api/attendance/admin/calendar', json=invalid)).status_code == 422
    assert (await client.post('/api/attendance/admin/calendar', json=dict(payload, kind='sick'))).status_code == 422


async def test_calendar_tenant_and_role_access(client, db_session, make_user, make_organization):
    org, employee, admin = await setup_employee(client, db_session, make_user, make_organization)
    response = await client.post('/api/attendance/admin/calendar', json=dict(employee_id=str(employee.id),
        date_from='2026-10-01', date_to='2026-10-01', kind='sick', reason='Approved leave'))
    assert response.status_code == 200
    row = await db_session.scalar(select(EmployeeCalendarDay).where(EmployeeCalendarDay.employee_id == employee.id))
    other = await make_organization(name='Other Workforce')
    user, password = await make_user(email='other-workforce@example.com', role=UserRole.org_admin, organization_id=other.id)
    await login(client, user.email, password)
    assert (await client.post(f'/api/attendance/admin/calendar/{row.id}/reset', json={'reason': 'Other tenant'})).status_code == 404
    assert (await client.post('/api/attendance/admin/calendar', json=dict(employee_id=str(employee.id),
        date_from='2026-10-01', date_to='2026-10-01', kind='off', reason='Other tenant'))).status_code == 404
    assert (await client.post(f'/api/attendance/admin/employees/{employee.id}/telegram')).status_code == 404
    assert (await client.get('/api/attendance/admin/calendar?month=2026-10')).json()['rows'] == []
    operator, password = await make_user(email='worker-workforce@example.com', role=UserRole.operator, organization_id=org.id)
    await login(client, operator.email, password)
    assert (await client.get('/api/attendance/admin/calendar?month=2026-10')).status_code == 403
    assert (await client.get('/api/attendance/admin/timesheet.xlsx?month=2026-10')).status_code == 403


async def test_report_breaks_and_incomplete_marks(client, db_session, make_user, make_organization):
    org, employee, admin = await setup_employee(client, db_session, make_user, make_organization)
    day = (utcnow() - timedelta(days=2)).date()
    for kind, hour in [('in', 9), ('out', 12), ('in', 13), ('out', 18)]:
        db_session.add(AttendanceEvent(organization_id=org.id, employee_id=employee.id, kind=kind,
            source='manual', occurred_at=datetime.combine(day, time(hour), tzinfo=timezone.utc)))
    await db_session.commit()
    report = await build_report(db_session, org, day, day)
    assert report['rows'][0]['worked_minutes'] == 480 and not report['rows'][0]['needs_review']
    events = (await db_session.scalars(select(AttendanceEvent).where(AttendanceEvent.employee_id == employee.id).order_by(AttendanceEvent.occurred_at))).all()
    await db_session.delete(events[-1]); await db_session.commit()
    report = await build_report(db_session, org, day, day)
    assert report['rows'][0]['worked_minutes'] == 180 and report['rows'][0]['needs_review']


async def test_late_reminder_deduplicated_and_absence_exempt(client, db_session, make_user, make_organization, monkeypatch):
    org, employee, admin = await setup_employee(client, db_session, make_user, make_organization)
    employee.telegram_chat_id = 98765
    await db_session.commit()
    monkeypatch.setattr(telegram, 'enabled', lambda: True)
    send = AsyncMock(return_value=True)
    monkeypatch.setattr(telegram, 'send_message', send)
    now = datetime(2026, 10, 6, 9, 10, tzinfo=timezone.utc)
    key = f'telegram:late:{employee.id}:2026-10-06:2026-10-06T09:00:00+00:00'
    try:
        await run_reminders_once(db_session, redis_client, now=now)
        await run_reminders_once(db_session, redis_client, now=now)
        assert send.await_count == 1
        await redis_client.delete(key)
        send.return_value = False
        await run_reminders_once(db_session, redis_client, now=now)
        assert not await redis_client.exists(key)  # Failed delivery may retry.
        send.return_value = True
        await run_reminders_once(db_session, redis_client, now=now)
        assert send.await_count == 3
        await redis_client.delete(key)
        db_session.add(EmployeeCalendarDay(organization_id=org.id, employee_id=employee.id, day=now.date(),
            kind='sick', reason='Approved absence', updated_by=admin.id))
        await db_session.commit()
        await run_reminders_once(db_session, redis_client, now=now)
        assert send.await_count == 3
    finally:
        await redis_client.delete(key)
